"""Unit tests for Dynamic 10-Year Excel Model Exporter."""

import os
from pathlib import Path
import openpyxl
import pytest
from src.export_excel_model import generate_10year_model_workbook


def test_excel_workbook_generation(tmp_path):
    out_file = tmp_path / "test_model.xlsx"
    path = generate_10year_model_workbook(
        client_name="Test Client LLC",
        starting_revenue=3000000.0,
        gross_margin=0.35,
        starting_cash=100000.0,
        output_path=str(out_file),
    )

    assert Path(path).exists()
    wb = openpyxl.load_workbook(path)
    assert "10-Year Growth Roadmap" in wb.sheetnames
    assert "13-Week Cash Flow War Room" in wb.sheetnames

    ws1 = wb["10-Year Growth Roadmap"]
    # Check that revenue exists in row 5
    assert ws1.cell(row=5, column=2).value == 3000000
